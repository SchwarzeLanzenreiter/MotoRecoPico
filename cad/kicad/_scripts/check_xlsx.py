# -*- coding: utf-8 -*-
"""Verify the recalculated workbook: no formula errors, the total is right."""
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from openpyxl import load_workbook

RECALC = os.path.join(os.environ['TEMP'], 'lorecalc', 'MotoRecoPico-部品購入先.xlsx')
ERRS = ('#REF!', '#NAME?', '#VALUE!', '#DIV/0!', '#N/A', '#NULL!', '#NUM!')

wb = load_workbook(RECALC, data_only=True)
bad = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.strip() in ERRS:
                print(f"ERROR {ws.title}!{c.coordinate} = {c.value}")
                bad += 1
print(f"formula errors: {bad}")

ws = wb['購入先一覧']
qty = [ws.cell(row=r, column=4).value for r in range(6, 29)]
qty = [q for q in qty if isinstance(q, int)]
print(f"rows with a quantity: {len(qty)}, python sum = {sum(qty)}")
for r in range(26, 32):
    v = ws.cell(row=r, column=4).value
    if v is not None and ws.cell(row=r, column=3).value == '合計':
        print(f"total cell D{r} (cached) = {v}")

wb2 = load_workbook(RECALC)
for r in range(26, 32):
    v = wb2['購入先一覧'].cell(row=r, column=4).value
    if isinstance(v, str) and v.startswith('='):
        print(f"formula at D{r}: {v}")

# hyperlinks survived?
n = sum(1 for ws in wb2.worksheets for row in ws.iter_rows()
        for c in row if c.hyperlink is not None)
print(f"hyperlinks: {n}")
